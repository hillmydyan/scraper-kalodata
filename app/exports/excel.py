"""
Excel exporter module using openpyxl.
Follows rules in PRD §10:
- Product_{COUNTRY}_run-{RUN_DATE}_data-{PERIOD}_{STATUS}.xlsx
- Sheet 'Data': Freeze pane, AutoFilter, column widths, currency formatting, ID as text, hyperlinks.
- Sheet 'Info': Run metadata, timestamps (WIB), rejection counts, version.
"""

import os
import hashlib
from datetime import datetime, date
from pathlib import Path
from typing import List, Optional, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import zoneinfo

from app.config import settings, get_country_config


WIB = zoneinfo.ZoneInfo("Asia/Jakarta")


class ExcelExporter:
    @staticmethod
    def format_wib(dt: Optional[datetime]) -> str:
        if dt is None:
            return ""
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=zoneinfo.ZoneInfo("UTC"))
        return dt.astimezone(WIB).strftime("%Y-%m-%d %H:%M:%S WIB")

    @classmethod
    def generate_excel_file(
        cls,
        country_code: str,
        run_date: date,
        period_date: date,
        status: str,
        items: List[Dict[str, Any]],
        job_metadata: Dict[str, Any],
        version: int = 1,
        module: str = "product"
    ) -> Optional[Dict[str, Any]]:
        """
        Generates Excel file per country and returns metadata dictionary.
        Supports both 'product' and 'creator' modules.
        """
        # If no data and export_empty_file is disabled
        if not items and status == "NO_DATA" and not settings.EXPORT_EMPTY_FILE:
            return None

        # Build filename
        v_suffix = f"_v{version}" if version > 1 else ""
        prefix = "Creator" if module == "creator" else "Product"
        file_name = f"{prefix}_{country_code.upper()}_run-{run_date}_data-{period_date}_{status}{v_suffix}.xlsx"

        # Directory structure runtime/exports/{YYYY}/{MM}/
        now = datetime.now(WIB)
        year_str = str(now.year)
        month_str = f"{now.month:02d}"
        export_dir = settings.RUNTIME_DIR / "exports" / year_str / month_str
        export_dir.mkdir(parents=True, exist_ok=True)
        target_path = export_dir / file_name

        wb = openpyxl.Workbook()
        
        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")

        # Country config for currency formatting
        c_cfg = get_country_config(country_code)
        currency_symbol = c_cfg.get("currency_symbol", "")
        if country_code == "ID":
            currency_format = f'"{currency_symbol} "#,##0'
        else:
            currency_format = f'"{currency_symbol} "#,##0.00'

        # 1. SHEET: Data
        ws_data = wb.active
        ws_data.title = "Data"
        ws_data.views.sheetView[0].showGridLines = True

        if module == "creator":
            headers = [
                "Nama Kreator",
                "Produk Ke",
                "ID Produk",
                "Nama Produk",
                "Nama Toko",
                "Pendapatan",
                "Penjualan",
                "Tingkat Komisi",
                "Link Produk",
                "Mata Uang",
                "Negara",
                "Periode Data",
                "Waktu Scraping",
                "Job ID"
            ]
        else:
            headers = [
                "ID Produk",
                "Nama Produk",
                "Nama Toko",
                "Link Produk",
                "Pendapatan",
                "Penjualan",
                "Tingkat Komisi",
                "Link Video",
                "Metode Pilih Video",
                "Mata Uang",
                "Negara",
                "Periode Data",
                "Waktu Scraping",
                "Job ID"
            ]

        ws_data.append(headers)

        # Style header row
        for col_num in range(1, len(headers) + 1):
            cell = ws_data.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Fill rows
        for row_idx, item in enumerate(items, start=2):
            prod_id = str(item.get("source_product_id") or "")
            prod_name = str(item.get("product_name") or "")
            shop_name = str(item.get("shop_name") or "")
            prod_url = str(item.get("product_url") or "")
            rev_val = item.get("revenue_value")
            sales_val = item.get("sales_value")
            comm_rate = item.get("commission_rate")
            video_url = str(item.get("video_url") or "")
            video_method = str(item.get("video_selection_method") or "")
            curr = str(item.get("currency_code") or country_code)
            country = str(item.get("country_code") or country_code)
            period_str = str(period_date)
            scraped_str = cls.format_wib(item.get("scraped_at") or datetime.utcnow())
            job_id = str(item.get("scrape_job_id") or "")

            if module == "creator":
                c_name = str(item.get("creator_name") or "")
                p_rank = item.get("product_rank")
                p_rank_str = f"Produk {p_rank}" if p_rank else ""
                row_data = [
                    c_name,
                    p_rank_str,
                    prod_id,
                    prod_name,
                    shop_name,
                    rev_val,
                    sales_val,
                    comm_rate,
                    prod_url,
                    curr,
                    country,
                    period_str,
                    scraped_str,
                    job_id
                ]
            else:
                row_data = [
                    prod_id,
                    prod_name,
                    shop_name,
                    prod_url,
                    rev_val,
                    sales_val,
                    comm_rate,
                    video_url,
                    video_method,
                    curr,
                    country,
                    period_str,
                    scraped_str,
                    job_id
                ]
            ws_data.append(row_data)

            # Apply cell types and formats
            if module == "creator":
                # Col 2: Produk Ke
                cell_prank = ws_data.cell(row=row_idx, column=2)
                cell_prank.alignment = Alignment(horizontal="center")

                # Col 3: ID as text
                cell_id = ws_data.cell(row=row_idx, column=3)
                cell_id.number_format = "@"
                cell_id.alignment = Alignment(horizontal="left")

                # Col 6: Pendapatan
                cell_rev = ws_data.cell(row=row_idx, column=6)
                if rev_val is not None:
                    cell_rev.number_format = currency_format

                # Col 7: Penjualan
                cell_sales = ws_data.cell(row=row_idx, column=7)
                if sales_val is not None:
                    cell_sales.number_format = "#,##0"

                # Col 8: Tingkat Komisi
                cell_comm = ws_data.cell(row=row_idx, column=8)
                if comm_rate is not None:
                    cell_comm.number_format = "0.00%"

                # Col 9: Link Produk
                if prod_url.startswith("http"):
                    cell_url = ws_data.cell(row=row_idx, column=9)
                    cell_url.hyperlink = prod_url
                    cell_url.font = link_font
            else:
                # Col 1: ID as text
                cell_id = ws_data.cell(row=row_idx, column=1)
                cell_id.number_format = "@"
                cell_id.alignment = Alignment(horizontal="left")

                # Col 4: Link Produk
                if prod_url.startswith("http"):
                    cell_url = ws_data.cell(row=row_idx, column=4)
                    cell_url.hyperlink = prod_url
                    cell_url.font = link_font

                # Col 5: Pendapatan
                cell_rev = ws_data.cell(row=row_idx, column=5)
                if rev_val is not None:
                    cell_rev.number_format = currency_format

                # Col 6: Penjualan
                cell_sales = ws_data.cell(row=row_idx, column=6)
                if sales_val is not None:
                    cell_sales.number_format = "#,##0"

                # Col 7: Tingkat Komisi
                cell_comm = ws_data.cell(row=row_idx, column=7)
                if comm_rate is not None:
                    cell_comm.number_format = "0.00%"

                # Col 8: Link Video
                if video_url.startswith("http"):
                    cell_vid = ws_data.cell(row=row_idx, column=8)
                    cell_vid.hyperlink = video_url
                    cell_vid.font = link_font

            # Borders for data row
            for col_num in range(1, len(headers) + 1):
                ws_data.cell(row=row_idx, column=col_num).border = border_thin

        # Freeze top row
        ws_data.freeze_panes = "A2"

        # Enable AutoFilter
        ws_data.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(items) + 1, 1)}"

        # Adjust column widths automatically
        for col in ws_data.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws_data.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # 2. SHEET: Info
        ws_info = wb.create_sheet(title="Info")
        ws_info.views.sheetView[0].showGridLines = True
        info_headers = ["Parameter", "Nilai"]
        ws_info.append(info_headers)
        for c in range(1, 3):
            cell = ws_info.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill

        info_rows = [
            ("Negara", country_code),
            ("Periode Data", str(period_date)),
            ("Tanggal Run", str(run_date)),
            ("Waktu Scraping (WIB)", cls.format_wib(datetime.utcnow())),
            ("Status Job", status),
            ("Metode Sumber Pendapatan", job_metadata.get("revenue_source_method", "product_filter")),
            ("Jumlah Produk Tersimpan", len(items)),
            ("Jumlah Baris Rejected", job_metadata.get("rows_rejected", 0)),
            ("Versi Ekspor", f"v{version}"),
            ("Versi Aplikasi", "2.0"),
            ("Catatan", job_metadata.get("notes", "Normal export"))
        ]

        for r_idx, (param, val) in enumerate(info_rows, start=2):
            ws_info.append([param, val])
            ws_info.cell(row=r_idx, column=1).font = Font(bold=True)
            ws_info.cell(row=r_idx, column=1).border = border_thin
            ws_info.cell(row=r_idx, column=2).border = border_thin

        ws_info.column_dimensions["A"].width = 28
        ws_info.column_dimensions["B"].width = 50

        # Save workbook
        wb.save(target_path)

        # Calculate file size and SHA256 checksum
        size_bytes = os.path.getsize(target_path)
        hasher = hashlib.sha256()
        with open(target_path, "rb") as f:
            while chunk := f.read(65536):
                hasher.update(chunk)
        checksum = hasher.hexdigest()

        rel_path = f"{year_str}/{month_str}/{file_name}"

        return {
            "file_name": file_name,
            "relative_path": rel_path,
            "absolute_path": target_path,
            "row_count": len(items),
            "size_bytes": size_bytes,
            "checksum_sha256": checksum,
            "status": "SUCCESS" if items else ("EMPTY" if status == "NO_DATA" else status)
        }
