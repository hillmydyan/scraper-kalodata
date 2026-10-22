"""
Service for Auto-Analysis and Opportunity Scoring (Revenue vs. Commission).
Provides smart insights to identify high-revenue, low-commission opportunities.
"""

from typing import List, Dict, Any, Optional
from datetime import date
from io import BytesIO
from sqlalchemy import select, desc
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models import SessionLocal
from app.models.entities import KalodataProductResult, KalodataCreatorProductResult, ScrapeJob


class AnalyticsService:

    @staticmethod
    def get_analyzed_items(
        db,
        country: Optional[str] = None,
        module: Optional[str] = "all",
        max_commission: float = 0.10,  # 10%
        min_revenue: float = 0.0,
        tier_filter: Optional[str] = None,
        search_query: Optional[str] = None,
        sort_by: str = "opportunity"
    ) -> Dict[str, Any]:
        """
        Retrieves, normalizes, scores, and categorizes products and creator-products.
        """
        raw_items: List[Dict[str, Any]] = []

        # 1. Fetch Product Results
        if module in ("all", "product"):
            stmt = select(KalodataProductResult)
            if country:
                stmt = stmt.filter(KalodataProductResult.country_code == country.upper())
            prod_records = db.execute(stmt).scalars().all()
            for p in prod_records:
                rev = float(p.revenue_value) if p.revenue_value is not None else 0.0
                comm = float(p.commission_rate) if p.commission_rate is not None else None
                sales = int(p.sales_value) if p.sales_value is not None else 0

                raw_items.append({
                    "id": p.id,
                    "source": "product",
                    "source_product_id": p.source_product_id,
                    "product_name": p.product_name or "-",
                    "shop_name": p.shop_name or "-",
                    "creator_name": None,
                    "country_code": p.country_code,
                    "period_start": p.period_start,
                    "currency_code": p.currency_code or "IDR",
                    "revenue_value": rev,
                    "revenue_raw": p.revenue_raw or "-",
                    "sales_value": sales,
                    "sales_raw": p.sales_raw or str(sales),
                    "commission_rate": comm,
                    "commission_raw": p.commission_raw or (f"{comm*100:.1f}%" if comm is not None else "-"),
                    "product_url": p.product_url,
                    "video_url": p.video_url,
                    "scraped_at": p.scraped_at
                })

        # 2. Fetch Creator Product Results
        if module in ("all", "creator"):
            stmt_c = select(KalodataCreatorProductResult)
            if country:
                stmt_c = stmt_c.filter(KalodataCreatorProductResult.country_code == country.upper())
            creator_records = db.execute(stmt_c).scalars().all()
            for c in creator_records:
                rev = float(c.revenue_value) if c.revenue_value is not None else 0.0
                comm = float(c.commission_rate) if c.commission_rate is not None else None
                sales = int(c.sales_value) if c.sales_value is not None else 0

                raw_items.append({
                    "id": c.id,
                    "source": "creator",
                    "source_product_id": c.source_product_id,
                    "product_name": c.product_name or "-",
                    "shop_name": c.shop_name or "-",
                    "creator_name": c.creator_name or "-",
                    "creator_rank": c.creator_rank,
                    "product_rank": c.product_rank,
                    "country_code": c.country_code,
                    "period_start": c.period_start,
                    "currency_code": c.currency_code or "IDR",
                    "revenue_value": rev,
                    "revenue_raw": c.revenue_raw or "-",
                    "sales_value": sales,
                    "sales_raw": c.sales_raw or str(sales),
                    "commission_rate": comm,
                    "commission_raw": c.commission_raw or (f"{comm*100:.1f}%" if comm is not None else "-"),
                    "product_url": c.product_url,
                    "video_url": None,
                    "scraped_at": c.scraped_at
                })

        # 3. Calculate Scoring and Tiering
        processed_candidates = []
        valid_revenues = [item["revenue_value"] for item in raw_items if item["revenue_value"] > 0]
        median_rev = 10_000_000.0  # default 10 juta
        if valid_revenues:
            sorted_rev = sorted(valid_revenues)
            median_rev = sorted_rev[len(sorted_rev) // 2]
            # Ensure minimum reasonable benchmark
            median_rev = max(median_rev, 5_000_000.0)

        for item in raw_items:
            rev = item["revenue_value"]
            comm = item["commission_rate"]  # decimal, e.g. 0.07 for 7%

            # Filter Search Keyword
            if search_query:
                sq = search_query.lower()
                matches = (
                    sq in item["product_name"].lower()
                    or sq in (item["shop_name"] or "").lower()
                    or sq in (item["creator_name"] or "").lower()
                    or sq in item["source_product_id"].lower()
                )
                if not matches:
                    continue

            # Calculate Opportunity Score:
            # Score = (Revenue / 1,000,000) / max(Commission_Pct, 0.5)
            # Only calculated for verified commission data to prioritize proven high-revenue low-commission winners.
            comm_pct = (comm * 100) if comm is not None else None
            if rev > 0 and comm_pct is not None and comm_pct > 0:
                opp_score = round((rev / 1_000_000.0) / max(comm_pct, 0.5), 2)
            else:
                opp_score = 0.0

            # Determine Tier (Sesuai aturan: Rendah <= 7%, Sedang 8-10%, Tinggi >= 11%)
            if comm_pct is not None:
                if comm <= 0.0701:  # 7% ke bawah
                    tier = "LOW_COMM"
                    tier_label = "Komisi Rendah"
                    tier_badge = "status-SUCCESS"
                elif comm <= 0.1001:  # 8% - 10%
                    tier = "MID_COMM"
                    tier_label = "Komisi Sedang"
                    tier_badge = "status-PARTIAL"
                else:  # 11% ke atas
                    tier = "HIGH_COMM"
                    tier_label = "Komisi Tinggi"
                    tier_badge = "status-FAILED"
            else:
                tier = "UNVERIFIED"
                tier_label = "Data Masked"
                tier_badge = "status-NO_DATA"

            item["commission_pct"] = comm_pct
            item["opportunity_score"] = opp_score
            item["tier"] = tier
            item["tier_label"] = tier_label
            item["tier_badge"] = tier_badge

            processed_candidates.append(item)

        # Champions: 1 produk dengan revenue tertinggi per kategori komisi
        low_candidates = [i for i in processed_candidates if i["tier"] == "LOW_COMM" and i["revenue_value"] > 0]
        mid_candidates = [i for i in processed_candidates if i["tier"] == "MID_COMM" and i["revenue_value"] > 0]
        high_candidates = [i for i in processed_candidates if i["tier"] == "HIGH_COMM" and i["revenue_value"] > 0]

        champion_low = max(low_candidates, key=lambda x: (x["revenue_value"], x["sales_value"])) if low_candidates else None
        champion_mid = max(mid_candidates, key=lambda x: (x["revenue_value"], x["sales_value"])) if mid_candidates else None
        champion_high = max(high_candidates, key=lambda x: (x["revenue_value"], x["sales_value"])) if high_candidates else None

        # Filter items for main table
        scored_items = []
        effective_tier_filter = tier_filter
        if sort_by == "rev_low_comm" and (not tier_filter or tier_filter == "ALL"):
            effective_tier_filter = "LOW_COMM"
        elif sort_by == "rev_mid_comm" and (not tier_filter or tier_filter == "ALL"):
            effective_tier_filter = "MID_COMM"
        elif sort_by == "rev_high_comm" and (not tier_filter or tier_filter == "ALL"):
            effective_tier_filter = "HIGH_COMM"

        for item in processed_candidates:
            rev = item["revenue_value"]
            tier = item["tier"]

            if min_revenue > 0 and rev < min_revenue:
                continue
            if effective_tier_filter and effective_tier_filter != "ALL" and tier != effective_tier_filter:
                continue

            scored_items.append(item)

        # 4. Sorting
        if sort_by in ("rev_low_comm", "rev_mid_comm", "rev_high_comm", "revenue"):
            # Sort by highest revenue / GMV first, then highest sales
            scored_items.sort(key=lambda x: (x["revenue_value"], x["sales_value"]), reverse=True)
        elif sort_by == "sales":
            # Sort by highest sales volume (pcs) first, then revenue
            scored_items.sort(key=lambda x: (x["sales_value"], x["revenue_value"]), reverse=True)
        elif sort_by == "commission_asc":
            # Lowest commission percentage first
            scored_items.sort(key=lambda x: (x["commission_pct"] is None, x["commission_pct"] or 999.0, -x["revenue_value"]))
        elif sort_by == "commission_desc":
            # Highest commission percentage first
            scored_items.sort(key=lambda x: (x["commission_pct"] is None, -(x["commission_pct"] or 0.0), -x["revenue_value"]))
        else:
            scored_items.sort(key=lambda x: (x["sales_value"], x["revenue_value"]), reverse=True)

        # 5. Calculate KPI Metrics
        total_items = len(scored_items)
        low_comm_items = [i for i in scored_items if i["tier"] == "LOW_COMM"]
        low_comm_count = len(low_comm_items)

        valid_comm_items = [i["commission_pct"] for i in scored_items if i["commission_pct"] is not None]
        avg_commission = round(sum(valid_comm_items) / len(valid_comm_items), 1) if valid_comm_items else 0.0

        max_revenue = max([i["revenue_value"] for i in scored_items], default=0.0)
        total_revenue_sum = sum([i["revenue_value"] for i in scored_items])

        total_sales_sum = sum([i["sales_value"] for i in scored_items])
        max_sales = max([i["sales_value"] for i in scored_items], default=0)
        top_selling_item = max(scored_items, key=lambda x: x["sales_value"]) if scored_items else None

        # Chart Data points (Revenue in Millions vs Commission Pct)
        scatter_points = []
        for i in scored_items[:80]:
            if i["commission_pct"] is not None and i["revenue_value"] > 0:
                scatter_points.append({
                    "x": round(i["commission_pct"], 1),
                    "y": round(i["revenue_value"] / 1_000_000.0, 2),
                    "name": i["product_name"][:35],
                    "shop": i["shop_name"] or "-",
                    "tier": i["tier"],
                    "sales": i["sales_value"]
                })

        tier_counts = {
            "LOW_COMM": len([i for i in scored_items if i["tier"] == "LOW_COMM"]),
            "MID_COMM": len([i for i in scored_items if i["tier"] == "MID_COMM"]),
            "HIGH_COMM": len([i for i in scored_items if i["tier"] == "HIGH_COMM"]),
            "UNVERIFIED": len([i for i in scored_items if i["tier"] == "UNVERIFIED"]),
        }

        return {
            "items": scored_items,
            "kpis": {
                "total_items": total_items,
                "low_comm_count": low_comm_count,
                "avg_commission": avg_commission,
                "max_revenue": max_revenue,
                "total_revenue_sum": total_revenue_sum,
                "total_sales_sum": total_sales_sum,
                "max_sales": max_sales,
                "top_selling_item": top_selling_item,
                "champion_low": champion_low,
                "champion_mid": champion_mid,
                "champion_high": champion_high,
                "median_benchmark": median_rev
            },
            "chart": {
                "scatter": scatter_points,
                "tier_counts": tier_counts
            }
        }

    @staticmethod
    def generate_excel_export(items: List[Dict[str, Any]]) -> BytesIO:
        """
        Exports analyzed items list to styled Excel.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Smart Analytics"

        # Headers
        headers = [
            "Peringkat", "Tingkat Komisi", "Penjualan (Qty)", "Nama Produk",
            "Nama Toko / Brand", "Kreator", "Pendapatan (Revenue GMV)", "Komisi (%)",
            "Negara", "Periode", "Tautan Produk", "Tautan Video"
        ]
        ws.append(headers)

        # Styling
        header_fill = PatternFill(start_color="784428", end_color="784428", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style="thin", color="EADFCE"),
            right=Side(style="thin", color="EADFCE"),
            top=Side(style="thin", color="EADFCE"),
            bottom=Side(style="thin", color="EADFCE")
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 2
        for rank, item in enumerate(items, 1):
            comm_str = f"{item['commission_pct']:.1f}%" if item.get("commission_pct") is not None else "-"
            tier_display = f"{comm_str} ({item.get('tier_label', '-')})" if item.get("commission_pct") is not None else item.get("tier_label", "-")
            rev_val = item.get("revenue_value", 0.0)

            ws.append([
                rank,
                tier_display,
                item.get("sales_value", 0),
                item.get("product_name", "-"),
                item.get("shop_name", "-"),
                item.get("creator_name") or "-",
                rev_val,
                comm_str,
                item.get("country_code", "-"),
                str(item.get("period_start", "-")),
                item.get("product_url") or "-",
                item.get("video_url") or "-"
            ])

            # Format Currency on column 7
            ws.cell(row=row_idx, column=7).number_format = '#,##0'

            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).border = border_thin

            row_idx += 1

            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).border = border_thin

            row_idx += 1

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
